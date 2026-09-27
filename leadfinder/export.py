"""Tabular views + Excel export (live profit formulas, same maths as profit_engine)."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ORDER = ["QUALIFIED", "REVIEW", "NEAR-MISS", "PACK MISMATCH", "FAILED BUY BOX VARIANCE",
         "NOT PROFITABLE", "MISSING FEE DATA", "NOT SOURCED", "PENDING"]


def to_frame(leads: list[dict]) -> pd.DataFrame:
    rows = []
    for l in leads:
        c = l.get("calc") or {}
        rows.append({
            "Status": l.get("status"), "Flag": l.get("flag") or "",
            "ASIN": l["asin"], "Title": l.get("title"), "Brand": l.get("brand"),
            "BSR": l.get("bsr"), "Monthly sold": l.get("monthly_sold"), "Offers": l.get("offers"),
            "Items": l.get("num_items") or l.get("package_qty"),
            "Buy Box £": l.get("buybox"), "BB 90d avg £": l.get("buybox_avg90"),
            "BB variance %": round(l["bb_variance_pct"], 1) if l.get("bb_variance_pct") is not None else None,
            "FBA fee £": l.get("fba_fee"), "Referral %": c.get("referral_pct_used") or l.get("referral_pct"),
            "Source": l.get("site") or "", "Source price £": l.get("source_price"), "VAT": l.get("vat") or "",
            "Cost inc VAT £": c.get("cost_inc_vat"), "Referral £": c.get("referral_fee"),
            "Digital £": c.get("digital_fee"), "Net VAT £": c.get("net_vat"),
            "Profit £": c.get("profit"), "ROI %": c.get("roi_pct"), "Margin %": c.get("margin_pct"),
            "Confidence": l.get("confidence") or "", "Pack match": l.get("pack_match") or "",
            "Method": l.get("method") or "", "Notes": l.get("notes") or "",
            "Source URL": l.get("url") or "", "Amazon": l.get("amazon_url"), "EAN": l.get("ean"),
        })
    df = pd.DataFrame(rows)
    front = ["Status", "Flag", "ASIN", "Title", "Profit £", "ROI %", "Buy Box £", "Source price £", "Source",
             "Items", "Pack match", "Confidence"]
    if not df.empty:
        df = df[front + [c for c in df.columns if c not in front]]
        df["_o"] = df["Status"].map({s: i for i, s in enumerate(ORDER)}).fillna(99)
        df = df.sort_values(["_o", "Profit £"], ascending=[True, False]).drop(columns="_o")
    return df


def to_excel(leads: list[dict], params: dict, summary: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    ws["A1"] = "Keepa lead run"; ws["A1"].font = Font(bold=True, size=14)
    r = 3
    for k, v in params.items():
        if k in ("keepa_query_override",) and not v:
            continue
        ws.cell(r, 1, k).font = bold
        ws.cell(r, 2, str(v))
        r += 1
    r += 1
    for k, v in (summary or {}).items():
        ws.cell(r, 1, k).font = bold
        ws.cell(r, 2, str(v))
        r += 1
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90

    heads = ["Status", "Flag", "ASIN", "Title", "Brand", "BSR", "Monthly sold", "Offers", "Items",
             "Buy Box £", "FBA fee £", "Referral %", "Source", "Source price £", "VAT",
             "Cost inc VAT £", "Referral £", "Digital £", "Net VAT £", "Profit £", "ROI %",
             "Confidence", "Pack match", "Notes", "Source URL", "Amazon"]
    fills = {"QUALIFIED": "C6EFCE", "REVIEW": "FFD966", "NEAR-MISS": "FFEB9C", "NOT SOURCED": "F2F2F2"}
    df = to_frame(leads)
    for name, subset in (("Qualified & review", df[df["Status"].isin(["QUALIFIED", "REVIEW"])] if not df.empty else df),
                         ("All leads", df)):
        s = wb.create_sheet(name)
        for i, h in enumerate(heads, 1):
            c = s.cell(1, i, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(wrap_text=True)
            s.column_dimensions[get_column_letter(i)].width = 45 if h in ("Title", "Notes") else 13
        s.freeze_panes = "D2"
        for ri, row in enumerate(subset.to_dict("records"), 2):
            vals = [row[h] if h in row else None for h in heads]
            for ci, v in enumerate(vals, 1):
                if isinstance(v, float) and pd.isna(v):
                    v = None
                s.cell(ri, ci, v)
            J, K, L, N, O, P, Q, R, S, T = (f"{c}{ri}" for c in "JKLNOPQRST")
            if row.get("Source price £") is not None and not pd.isna(row.get("Source price £")):
                s[P] = f'=IF({O}="ex",{N}*1.2,{N})'
                s[Q] = f'=MAX({J}*{L}/100,0.3)'
                s[R] = f'=ROUND({J}*0.007,2)'
                s[S] = f'=({J}-{P})/6'
                s[T] = f'={J}-{P}-{Q}-{K}-{R}-{S}'
                s[f"U{ri}"] = f'=IF({P}=0,"",{T}/{P}*100)'
            color = fills.get(row["Status"])
            if color:
                for ci in range(1, len(heads) + 1):
                    s.cell(ri, ci).fill = PatternFill("solid", fgColor=color)
        s.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{max(2, len(subset) + 1)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
